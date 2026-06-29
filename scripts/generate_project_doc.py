from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "LLM-Guardian_项目说明文档.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_run_style(run, size: int = 11, bold: bool = False, color: str | None = None) -> None:
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_paragraph(doc: Document, text: str, *, size: int = 11, bold: bool = False, space_after: int = 6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.45
    run = p.add_run(text)
    set_run_style(run, size=size, bold=bold)
    return p


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.line_spacing = 1.4
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(item)
        set_run_style(run, size=11)


def add_heading(doc: Document, text: str, level: int) -> None:
    heading = doc.add_paragraph()
    heading.style = f"Heading {level}"
    heading.paragraph_format.space_before = Pt(8 if level == 1 else 4)
    heading.paragraph_format.space_after = Pt(6)
    run = heading.add_run(text)
    set_run_style(run, size=16 if level == 1 else 13, bold=True, color="1F4E78" if level == 1 else "203864")


def build_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.3)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    run = title.add_run("LLM-Guardian 项目说明文档")
    set_run_style(run, size=20, bold=True, color="1F3A5F")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    run = subtitle.add_run(
        f"面向企业本地大模型服务的质量监测、SLA 审计与平台级 AI 诊断平台\n编制日期：{datetime.now().strftime('%Y-%m-%d')}"
    )
    set_run_style(run, size=10, color="5B6573")

    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(10)
    intro.paragraph_format.line_spacing = 1.45
    intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = intro.add_run(
        "本文档基于当前项目代码与已实现能力整理，重点说明平台建设背景、整体架构、核心功能边界及后续建设方向，"
        "便于管理层快速理解项目定位、实际价值与推广意义。"
    )
    set_run_style(run, size=11)

    add_heading(doc, "一、项目背景与目标", 1)
    add_paragraph(
        doc,
        "随着大模型在政企场景中的本地化部署持续增多，平台建设重点已经从“模型是否能够接入”，转向“模型上线后是否稳定、可控、可审计”。"
        "当前很多 AI 网关或接入层更关注协议兼容、统一鉴权和流量转发，但对模型服务质量本身缺少持续主动拨测与运营治理能力，"
        "因此很难第一时间发现服务退化、定位故障原因并形成可汇报的管理视图。",
    )
    add_paragraph(
        doc,
        "LLM-Guardian 正是在这一背景下建设的。项目聚焦企业本地大模型服务治理，围绕“持续观测、实时告警、周期审计、智能辅助诊断”形成闭环，"
        "支撑运维团队更好地落实中国电信集团“1 分钟发现、5 分钟定位、10 分钟解决”的要求。",
    )
    add_bullets(
        doc,
        [
            "建设目标 1：统一纳管 OpenAI 兼容接口、Ollama 兼容接口及自定义兼容接口，形成统一的模型服务台账与观测入口。",
            "建设目标 2：通过定时主动拨测持续采集 TTFT、TPS、总时延、成功率等关键指标，将零散调用体验沉淀为可追溯的数据资产。",
            "建设目标 3：通过告警、日志、SLA 审计和报告推送，构建从异常发现到管理汇报的闭环机制。",
            "建设目标 4：引入平台级第三方 AI 分析入口，结合模型部署信息与近期性能数据，定期生成资源健康诊断建议，辅助运维优化生产部署。",
            "建设目标 5：降低人工巡检、人工写报告和人工排障成本，为企业大模型平台提供可运营、可治理的质量保障底座。",
        ],
    )

    add_heading(doc, "二、项目整体架构及技术方案", 1)
    add_heading(doc, "2.1 整体架构", 2)
    add_paragraph(
        doc,
        "平台整体采用“前端控制台 + 后端服务 + 本地数据存储 + 定时调度 + 外部消息/模型接口”的轻量化架构，既便于单机快速部署演示，"
        "也具备向生产环境平滑演进的基础。整体业务链路如下：",
    )
    add_bullets(
        doc,
        [
            "模型纳管层：维护被监控模型的访问地址、密钥、模型标识、接口类型及本地部署信息。",
            "任务调度层：按设定周期发起主动拨测，并独立调度日报、周报、月报对应的 AI 分析任务。",
            "数据采集层：在每次拨测中记录请求结果、首字延迟、吞吐率、总时延、响应摘要、异常信息与 SLA 违约状态。",
            "分析治理层：对拨测日志进行聚合统计，生成实时看板、时序日志、SLA 合规审计和平台 AI 资源健康诊断建议。",
            "通知输出层：支持告警通知、PDF 报告导出与第三方消息推送，满足运维值守和管理汇报双重需要。",
        ],
    )

    add_heading(doc, "2.2 技术选型", 2)
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    for idx, text in enumerate(["层次", "技术选型", "说明"]):
        header_cells[idx].text = text
        set_cell_shading(header_cells[idx], "D9EAF7")
        for paragraph in header_cells[idx].paragraphs:
            for run in paragraph.runs:
                set_run_style(run, size=10, bold=True)

    rows = [
        ("前端", "React 18 + TypeScript + Vite + Tailwind CSS", "构建单页控制台，负责页面交互、图表展示、配置管理和报表操作。"),
        ("图表与视觉", "Recharts + Lucide React", "展示时序指标、SLA 统计、渠道状态和审计结果。"),
        ("后端", "FastAPI + SQLAlchemy", "提供 REST API、业务编排、数据持久化和权限控制。"),
        ("数据库", "SQLite", "当前版本采用轻量本地存储，适合单机部署、PoC 验证和中小规模演示。"),
        ("调度", "APScheduler", "统一调度拨测任务与日报、周报、月报三类 AI 分析缓存任务。"),
        ("网络调用", "aiohttp", "对接被监控模型接口、第三方 AI 分析接口及消息推送接口。"),
        ("外部集成", "OpenAI 兼容模型接口、飞书/钉钉/邮箱等消息通道", "满足模型拨测、平台级 AI 诊断和报告分发需求。"),
    ]
    for left, mid, right in rows:
        row_cells = table.add_row().cells
        for index, value in enumerate([left, mid, right]):
            row_cells[index].text = value
            for paragraph in row_cells[index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    set_run_style(run, size=10)

    add_heading(doc, "2.3 核心技术方案", 2)
    add_bullets(
        doc,
        [
            "统一适配方案：后端按接口类型选择不同探测适配器，兼容 OpenAI 类接口、Ollama 接口及自定义兼容接口，降低异构模型接入门槛。",
            "主动拨测方案：每个渠道可配置独立任务、探测频率、并发数、Prompt 及 SLA 阈值，系统自动执行拨测并沉淀原始日志。",
            "SLA 审计方案：围绕 TTFT、TPS、总时延、成功率、达标率进行聚合，形成日报、周报、月报的审计视图与报告输出。",
            "平台级 AI 诊断方案：不再由被监控模型自身生成建议，而是由单独配置的第三方 OpenAI 兼容分析模型统一生成并缓存诊断结果。",
            "兜底策略方案：当 AI 建议生成失败时，系统自动退回规则模板摘要，保证 SLA 合规页与报告中始终有稳定、可阅读的结果输出。",
            "消息联动方案：告警事件与周期报告可推送到飞书、钉钉、邮箱等渠道，提升异常触达与管理汇报效率。",
        ],
    )

    add_heading(doc, "2.4 当前架构边界说明", 2)
    add_bullets(
        doc,
        [
            "平台当前定位是大模型服务质量监测与治理平台，不替代 AI 网关的统一鉴权、流量分发、配额管理和多租户路由职责。",
            "平台级 AI 诊断入口当前首版采用单一 OpenAI 兼容模型配置，适合先行落地，后续可演进为多模型主备或分级分析体系。",
            "当前数据存储为 SQLite，更适合单节点或演示环境；若进入正式生产，建议升级至 PostgreSQL/MySQL 并补充备份与高可用方案。",
        ],
    )

    add_heading(doc, "三、功能模块说明", 1)
    modules = [
        (
            "3.1 渠道管理与模型纳管",
            [
                "平台支持新增、编辑、删除被监控模型渠道，维护名称、接口地址、密钥、模型标识、接口类型等基础信息。",
                "在同一 URL、同一密钥下可先获取模型列表，再选择具体纳管模型，降低多模型场景下的人工录入成本。",
                "对于本地私有化模型，平台可同时维护部署方式、部署配置、环境变量、部署参数等部署信息，为后续 AI 诊断提供上下文。",
            ],
            "适用场景：集团内部 AI 平台、本地私有化推理服务、第三方模型服务的统一纳管。",
            "实际价值：把原本分散在多个环境、多个地址中的模型服务统一纳入治理台账，为后续监测、审计和分析打下基础。",
        ),
        (
            "3.2 拨测配置与主动探测",
            [
                "每个纳管渠道均可配置独立拨测任务，包含探测文本、执行周期、并发数及指标阈值。",
                "平台自动调用目标模型接口，采集首字延迟、吞吐率、总时延、响应成功率、错误信息等关键数据。",
                "当任务保存或启用后，可触发即时拨测，帮助运维快速验证配置是否可用。",
            ],
            "适用场景：日常巡检、上线验收、变更后回归、容量观察、模型可用性验证。",
            "实际价值：从“被动等待用户报障”转变为“平台主动发现问题”，显著提升故障发现速度。",
        ),
        (
            "3.3 实时监控大盘",
            [
                "以可视化方式展示平台当前的拨测总量、成功率、SLA 水位、关键指标趋势和渠道健康状态。",
                "可帮助管理和运维快速把握当前整体运行情况，识别风险渠道与异常波动。",
                "页面定位偏运营总览，侧重整体健康态势与重点风险呈现。",
            ],
            "适用场景：值班巡检、领导参观演示、日常态势感知。",
            "实际价值：把复杂的性能数据转换成直观的管理视图，降低跨团队沟通成本。",
        ),
        (
            "3.4 时序测试日志与问题定位",
            [
                "平台保留每次拨测的时序日志，可按渠道、状态、时间范围等条件检索。",
                "日志不仅记录成功或失败，还保留性能细节、SLA 违约标识、异常摘要和响应片段，便于复盘。",
                "运维可围绕最近 2 小时、24 小时、一周、一个月等常用时间窗口快速定位问题集中区间。",
            ],
            "适用场景：异常排查、故障复盘、容量波动分析、模型稳定性核查。",
            "实际价值：支撑“5 分钟定位”，把问题从模糊感知转化为基于日志和指标的精准定位。",
        ),
        (
            "3.5 SLA 合规审计与周期报告",
            [
                "系统基于日报、周报、月报三类时间窗口，对模型渠道进行聚合评估，形成成功率、SLA 达标率、TPS、TTFT、ITL、总时延等统计结果。",
                "审计页强调吞吐率优先展示，体现对 TPS 的重点关注，便于快速识别模型服务是否具备持续输出能力。",
                "支持导出 PDF/CSV，并向飞书等第三方通道推送报告，方便管理汇报与留档审计。",
            ],
            "适用场景：周报月报汇报、服务承诺考核、跨部门沟通、客户侧服务复盘。",
            "实际价值：让 SLA 管理从技术视角扩展到管理视角，形成“可看、可讲、可汇报”的运营成果。",
        ),
        (
            "3.6 平台级 AI 资源健康诊断",
            [
                "平台新增独立的 AI 诊断配置入口，由管理员配置第三方 OpenAI 兼容分析模型，例如 DeepSeek 或 SiliconFlow 兼容接口。",
                "系统按日报、周报、月报周期自动汇总各渠道近期性能数据，并结合渠道部署信息生成中文资源健康诊断建议。",
                "该能力与被监控模型解耦，意味着即使某个被监控模型自身不可用，只要历史拨测数据存在，平台仍可对其形成分析结果。",
                "当第三方 AI 分析模型调用异常时，系统不再展示失败记录，而是自动使用规则模板生成性能摘要，保障页面和报告稳定可读。",
            ],
            "适用场景：领导查看诊断结论、运维做阶段性巡检、模型部署优化分析、资源扩缩容评估。",
            "实际价值：把“指标展示”升级为“指标解释”，帮助运维更快理解性能问题，并基于部署信息给出可执行的优化方向。",
        ),
        (
            "3.7 告警联动与消息推送",
            [
                "当拨测失败或关键指标超阈值时，平台可生成告警事件，并支持异常恢复通知。",
                "告警渠道支持飞书、钉钉、邮箱等方式，满足不同组织的通知偏好。",
                "除实时告警外，日报、周报、月报也可一键推送，并附带平台 AI 资源健康诊断建议，方便在群内直接查看摘要。",
            ],
            "适用场景：7×24 值班、群组协同、管理触达、重大异常升级。",
            "实际价值：实现“1 分钟发现”，缩短问题从发生到被人知晓的时间差。",
        ),
        (
            "3.8 系统设置与权限控制",
            [
                "平台支持本地管理员登录，核心配置项集中在系统设置中管理，包括 AI 诊断配置、消息推送相关设置等。",
                "前后端接口基于统一数据模型与权限判断进行处理，减少误操作风险。",
                "对于演示和私有化部署场景，当前方案轻量、直接，便于快速上线和维护。",
            ],
            "适用场景：内部演示、单点部署、平台管理员配置维护。",
            "实际价值：在不引入复杂基础设施的前提下，实现对关键能力的集中管理与受控使用。",
        ),
    ]
    for title_text, bullets, scenario, value in modules:
        add_heading(doc, title_text, 2)
        add_bullets(doc, bullets)
        add_paragraph(doc, scenario, bold=True, size=10, space_after=2)
        add_paragraph(doc, value, size=10)

    add_heading(doc, "四、应用场景与实际价值总结", 1)
    add_bullets(
        doc,
        [
            "面向集团内部 AI 平台：支撑本地大模型上线后的持续监测、SLA 管理和运维诊断。",
            "面向政企客户私有化交付：帮助客户建立可视化、可审计、可汇报的模型服务治理能力。",
            "面向运维团队：通过拨测、告警、日志、AI 诊断四类能力联动，降低巡检和排障成本。",
            "面向管理层：通过日报、周报、月报和飞书推送机制，快速掌握整体服务水位、重点风险和优化方向。",
        ],
    )
    add_paragraph(
        doc,
        "综合来看，平台的核心价值不在于“把模型接进来”，而在于让大模型服务在生产环境中真正做到可观测、可审计、可告警、可诊断、可汇报，"
        "从而推动企业大模型能力从“可用”迈向“可运营、可治理”。",
    )

    add_heading(doc, "五、后续优化规划与建设建议", 1)
    add_heading(doc, "5.1 能力优化方向", 2)
    add_bullets(
        doc,
        [
            "增强多模型分析能力：从单一平台级 AI 分析模型扩展为多模型主备、分场景分析或成本优先策略。",
            "完善报告中心：支持更丰富的模板化报告、管理驾驶舱摘要和长期趋势对比分析。",
            "增强异常根因分析：逐步结合 GPU、容器、节点、网关等基础设施指标，形成更完整的根因定位链路。",
            "补充租户与组织能力：适配更复杂的部门、客户、项目级管理场景，提升平台横向推广能力。",
        ],
    )
    add_heading(doc, "5.2 架构演进建议", 2)
    add_bullets(
        doc,
        [
            "数据库升级：正式生产建议由 SQLite 升级到 PostgreSQL 或 MySQL，并建立备份、恢复和审计机制。",
            "服务高可用：推动前后端容器化部署，结合反向代理、健康检查和进程守护提升稳定性。",
            "消息与调度解耦：后续可引入消息队列或任务中心，支撑更大规模的拨测与报告生成需求。",
            "安全加固：补充更细颗粒度的权限控制、密钥托管、访问审计与配置脱敏能力。",
        ],
    )
    add_heading(doc, "5.3 建设落地建议", 2)
    add_bullets(
        doc,
        [
            "建议先在集团内部或重点客户场景中形成示范点，优先验证本地模型纳管、主动拨测、AI 诊断和报告推送闭环。",
            "建议与 GPU 容器平台、AI 网关、传统监控平台形成协同，而非替代关系，共同构建大模型服务治理体系。",
            "建议围绕中国电信集团“1-5-10”要求设定量化考核指标，持续验证平台在发现、定位、辅助处置上的实际效果。",
        ],
    )

    doc.add_section(WD_SECTION.CONTINUOUS)
    footer = doc.sections[-1].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("LLM-Guardian 项目说明文档")
    set_run_style(run, size=9, color="7A7A7A")
    return doc


if __name__ == "__main__":
    document = build_document()
    document.save(OUTPUT)
    print(OUTPUT)
